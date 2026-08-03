import re

def parse_questions(text: str) -> list[dict]:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    
    # 1. Primary check: If text is separated by +++++ (question dividers)
    if re.search(r'\+{3,}', text):
        raw_blocks = re.split(r'\n?\s*\+{3,}\s*\n?', text)
    # 2. Secondary check: Split by empty lines
    elif '\n\n' in text:
        raw_blocks = re.split(r'\n\s*\n+', text)
    # 3. Fallback: Split by numbered lines (1., 2., Q1.)
    else:
        raw_blocks = re.split(r'\n(?=(?:\d+|[Qq]\d+)[\.\)\-]\s+)', text)

    questions = []
    for block in raw_blocks:
        block = block.strip()
        if not block:
            continue
        
        # If a block itself contains multiple questions separated by numbers e.g. "1. Q1 \n 2. Q2"
        sub_blocks = re.split(r'\n(?=(?:\d+|[Qq]\d+)[\.\)\-]\s+)', block)
        for sb in sub_blocks:
            sb = sb.strip()
            if not sb:
                continue
            q = _parse_block(sb)
            if q:
                questions.append(q)
                
    return questions

def _parse_block(block: str) -> dict | None:
    lines = [l.strip() for l in block.splitlines() if l.strip()]
    if not lines:
        return None

    # Check if lines have internal dividers like ===== or ===
    if any(re.match(r'^={3,}$', l) for l in lines):
        return _parse_equal_separated_block(lines)

    # Check if lines use A), B), C), D) or #A), #B)
    has_explicit_letters = any(re.match(r'^(?:#|\+|\*)?\s*([a-dA-D])\s*[\)\.]', l) for l in lines)
    if has_explicit_letters:
        return _parse_letter_options(lines)

    # Vertical options where options are listed line by line (with optional # on correct answer)
    return _parse_vertical_options(lines)

def _parse_equal_separated_block(lines: list[str]) -> dict | None:
    parts = []
    current_part = []
    for line in lines:
        if re.match(r'^={3,}$', line):
            if current_part:
                parts.append(" ".join(current_part).strip())
                current_part = []
        else:
            current_part.append(line)
    if current_part:
        parts.append(" ".join(current_part).strip())

    parts = [p for p in parts if p]
    if len(parts) < 2:
        return None

    q_text = parts[0]
    q_text = re.sub(r'^(?:\d+|[Qq]\d+)[\.\)\-]?\s*', '', q_text).strip()

    raw_options = parts[1:]
    letters = ["A", "B", "C", "D", "E"]
    variants = {}
    correct_idx = 0

    for idx, opt in enumerate(raw_options):
        if idx >= len(letters):
            break
        if opt.startswith("#") or opt.startswith("+") or opt.startswith("*"):
            correct_idx = idx
            opt = opt.lstrip("#+* ").strip()
        variants[letters[idx]] = opt

    if len(variants) < 2:
        return None

    correct_letter = letters[correct_idx] if 0 <= correct_idx < len(letters) else "A"
    return {
        "text": q_text,
        "variant_a": variants.get("A", ""),
        "variant_b": variants.get("B", ""),
        "variant_c": variants.get("C", ""),
        "variant_d": variants.get("D", ""),
        "correct_answer": correct_letter,
    }

def _parse_letter_options(lines: list[str]) -> dict | None:
    question_lines = []
    variants = {}
    correct_answer = None
    current_variant = None

    for line in lines:
        m_correct = re.match(r'^(?:#|===+|\+|\*)\s*([a-dA-D])\s*[\)\.]?\s*(.*)', line)
        if not m_correct:
            m_correct = re.match(r'^([a-dA-D])\s*[\)\.]?\s*[\*\+]\s*(.*)', line)

        if m_correct:
            letter = m_correct.group(1).upper()
            val = m_correct.group(2).strip()
            correct_answer = letter
            variants[letter] = val
            current_variant = letter
            continue

        m_normal = re.match(r'^([a-dA-D])\s*[\)\.]\s*(.*)', line)
        if m_normal:
            letter = m_normal.group(1).upper()
            val = m_normal.group(2).strip()
            if letter not in variants:
                variants[letter] = val
            current_variant = letter
            continue

        if current_variant and current_variant in variants:
            variants[current_variant] = (variants[current_variant] + " " + line).strip()
        else:
            cleaned = re.sub(r'^(?:\d+|[Qq]\d+)[\.\)\-]?\s*', '', line) if not question_lines else line
            question_lines.append(cleaned)

    if not correct_answer and len(variants) >= 2:
        correct_answer = "A"

    if len(variants) < 2:
        return None

    q_txt = " ".join(question_lines).strip()
    return {
        "text": q_txt,
        "variant_a": variants.get("A", ""),
        "variant_b": variants.get("B", ""),
        "variant_c": variants.get("C", ""),
        "variant_d": variants.get("D", ""),
        "correct_answer": correct_answer,
    }

def _parse_vertical_options(lines: list[str]) -> dict | None:
    if len(lines) < 3:
        return None

    opt_start_idx = -1
    for i, l in enumerate(lines):
        if l.startswith("#") or l.startswith("+") or l.startswith("*"):
            opt_start_idx = i
            break

    letters = ["A", "B", "C", "D", "E"]
    raw_options = []
    correct_idx = 0

    if opt_start_idx != -1:
        if len(lines) >= 4 and opt_start_idx >= 2:
            q_lines = lines[:-4] if len(lines) >= 5 else lines[:1]
            opt_lines = lines[len(q_lines):]
        else:
            q_lines = lines[:opt_start_idx]
            opt_lines = lines[opt_start_idx:]

        question_lines = [re.sub(r'^(?:\d+|[Qq]\d+)[\.\)\-]?\s*', '', l) for l in q_lines]
        for line in opt_lines:
            if line.startswith("#") or line.startswith("+") or line.startswith("*"):
                correct_idx = len(raw_options)
            clean_opt = line.lstrip("#+* ").strip()
            raw_options.append(clean_opt)
    else:
        question_lines = [re.sub(r'^(?:\d+|[Qq]\d+)[\.\)\-]?\s*', '', lines[0])]
        for line in lines[1:]:
            raw_options.append(line.strip())

    if len(raw_options) < 2:
        return None

    variants = {}
    correct_letter = letters[correct_idx] if 0 <= correct_idx < len(letters) else "A"
    for idx, opt in enumerate(raw_options):
        if idx >= len(letters):
            break
        variants[letters[idx]] = opt

    return {
        "text": " ".join(question_lines).strip(),
        "variant_a": variants.get("A", ""),
        "variant_b": variants.get("B", ""),
        "variant_c": variants.get("C", ""),
        "variant_d": variants.get("D", ""),
        "correct_answer": correct_letter,
    }

def parse_docx(file_path: str) -> str:
    from docx import Document
    doc = Document(file_path)
    lines = []
    for p in doc.paragraphs:
        if p.text and p.text.strip():
            lines.append(p.text.strip())
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text and cell.text.strip():
                    lines.append(cell.text.strip())
    return "\n".join(lines)

def parse_txt(file_path: str) -> str:
    with open(file_path, "r", encoding="utf-8") as f:
        return f.read()